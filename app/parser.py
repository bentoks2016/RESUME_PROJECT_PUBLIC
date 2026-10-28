from io import BytesIO
from pathlib import Path

import pymupdf
import pytesseract

from PIL import Image
from docx import Document


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
}


def ocr_page(page) -> str:
    """
    Render a PDF page at higher resolution and
    extract text using Tesseract OCR.
    """

    matrix = pymupdf.Matrix(
        2.5,
        2.5
    )

    pixmap = page.get_pixmap(
        matrix=matrix,
        alpha=False
    )

    image = Image.open(
        BytesIO(
            pixmap.tobytes("png")
        )
    )

    text = pytesseract.image_to_string(
        image,
        lang="eng"
    )

    return text.strip()


def extract_pdf(content: bytes) -> str:
    """
    Extract text from a PDF.

    Normal text extraction is attempted first.

    Pages without a useful text layer automatically
    fall back to OCR.
    """

    document = pymupdf.open(
        stream=content,
        filetype="pdf"
    )

    extracted_pages = []

    try:

        print(
            f"PDF contains {len(document)} page(s)."
        )

        for number, page in enumerate(
            document,
            start=1
        ):

            text = page.get_text(
                "text"
            ).strip()

            print(
                f"PDF page {number}: "
                f"{len(text)} characters "
                "from normal extraction."
            )

            # There is a useful normal text layer.
            if len(text) >= 40:

                extracted_pages.append(
                    text
                )

                continue

            print(
                f"PDF page {number}: "
                "attempting OCR..."
            )

            try:

                ocr_text = ocr_page(
                    page
                )

                print(
                    f"PDF page {number}: "
                    f"{len(ocr_text)} characters "
                    "from OCR."
                )

                if ocr_text:
                    extracted_pages.append(
                        ocr_text
                    )

            except Exception as error:

                print(
                    f"OCR ERROR on page "
                    f"{number}: {error}"
                )

    finally:

        document.close()

    result = "\n\n".join(
        extracted_pages
    ).strip()

    print(
        "Total PDF characters extracted:",
        len(result)
    )

    return result


def extract_docx(content: bytes) -> str:
    """
    Extract text from DOCX paragraphs and tables.
    """

    document = Document(
        BytesIO(content)
    )

    pieces = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            pieces.append(
                text
            )

    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:

                    cells.append(
                        text
                    )

            if cells:

                pieces.append(
                    " | ".join(cells)
                )

    result = "\n".join(
        pieces
    ).strip()

    print(
        "Total DOCX characters extracted:",
        len(result)
    )

    return result


def extract_txt(content: bytes) -> str:

    text = content.decode(
        "utf-8",
        errors="ignore"
    ).strip()

    print(
        "Total TXT characters extracted:",
        len(text)
    )

    return text


def extract_text(
    filename: str,
    content: bytes
) -> str:

    extension = Path(
        filename
    ).suffix.lower()

    print()
    print("=" * 60)
    print("PARSING FILE")
    print("Filename:", filename)
    print("Extension:", extension)
    print("File size:", len(content), "bytes")
    print("=" * 60)

    if extension not in ALLOWED_EXTENSIONS:

        raise ValueError(
            "Unsupported file type. "
            "Please upload PDF, DOCX, or TXT."
        )

    if extension == ".pdf":

        return extract_pdf(
            content
        )

    if extension == ".docx":

        return extract_docx(
            content
        )

    return extract_txt(
        content
    )


# =========================================================
# CONTACT INFORMATION EXTRACTION
# =========================================================

import re


def extract_contact_information(text: str) -> dict:
    """
    Deterministically extract contact information from
    source resume text.

    These values should be preserved rather than rewritten
    or invented by the LLM.
    """

    contact = {
        "email": "",
        "phone": "",
        "linkedin": "",
        "location": "",
    }


    # EMAIL
    email_match = re.search(
        r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b',
        text
    )

    if email_match:
        contact["email"] = email_match.group(0)


    # LINKEDIN
    linkedin_match = re.search(
        r'(?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9_%\-./]+',
        text,
        re.IGNORECASE
    )

    if linkedin_match:
        contact["linkedin"] = linkedin_match.group(0).rstrip(
            ".,;)"
        )


    # PHONE
    phone_patterns = [
        r'\+\d{1,3}[\s\-().]*\d{3,}[\d\s\-().]{5,}',
        r'\b\d{3}[\s.-]\d{3}[\s.-]\d{4}\b',
        r'\b\d{10}\b',
    ]

    for pattern in phone_patterns:

        phone_match = re.search(
            pattern,
            text
        )

        if phone_match:

            contact["phone"] = (
                phone_match.group(0).strip()
            )

            break


    return contact
