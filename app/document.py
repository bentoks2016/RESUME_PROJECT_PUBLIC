from io import BytesIO

from docx import Document

from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_TAB_ALIGNMENT,
)

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx.shared import (
    Inches,
    Pt,
)

from app.models import (
    CoverLetterDocument,
    ResumeDocument,
)


# =========================================================
# SHARED DOCUMENT CONFIGURATION
# =========================================================

def configure_document(
    document: Document
):

    section = document.sections[0]

    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.45)

    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = document.styles["Normal"]

    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0


def add_top_border(
    paragraph,
    size="16",
    color="000000"
):

    p = paragraph._p

    pPr = p.get_or_add_pPr()

    pBdr = OxmlElement(
        "w:pBdr"
    )

    top = OxmlElement(
        "w:top"
    )

    top.set(
        qn("w:val"),
        "single"
    )

    top.set(
        qn("w:sz"),
        size
    )

    top.set(
        qn("w:space"),
        "2"
    )

    top.set(
        qn("w:color"),
        color
    )

    pBdr.append(
        top
    )

    pPr.append(
        pBdr
    )


def add_bottom_border(
    paragraph,
    size="4",
    color="808080"
):

    p = paragraph._p

    pPr = p.get_or_add_pPr()

    pBdr = OxmlElement(
        "w:pBdr"
    )

    bottom = OxmlElement(
        "w:bottom"
    )

    bottom.set(
        qn("w:val"),
        "single"
    )

    bottom.set(
        qn("w:sz"),
        size
    )

    bottom.set(
        qn("w:space"),
        "2"
    )

    bottom.set(
        qn("w:color"),
        color
    )

    pBdr.append(
        bottom
    )

    pPr.append(
        pBdr
    )


def add_candidate_header(
    document,
    name,
    contact
):
    """
    Shared header used by both the resume and cover letter.
    """

    # Name
    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(
        name
    )

    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(15)


    # Heavy line
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(1)

    add_top_border(
        paragraph
    )


    # Phone | Email | LinkedIn
    contact_parts = []

    if contact.phone:
        contact_parts.append(
            contact.phone
        )

    if contact.email:
        contact_parts.append(
            contact.email
        )

    if contact.linkedin:
        contact_parts.append(
            contact.linkedin
        )


    if contact_parts:

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.space_after = Pt(0)

        for index, item in enumerate(
            contact_parts
        ):

            if index > 0:

                separator = paragraph.add_run(
                    " | "
                )

                separator.font.name = "Arial"
                separator.font.size = Pt(10)


            run = paragraph.add_run(
                item
            )

            run.font.name = "Arial"
            run.font.size = Pt(10)


    # City / Province or other source location
    if contact.location:

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        paragraph.paragraph_format.space_after = Pt(5)

        run = paragraph.add_run(
            contact.location
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)


def add_section_heading(
    document,
    title
):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(4)

    run = paragraph.add_run(
        title.upper()
    )

    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)

    add_bottom_border(
        paragraph
    )


def add_bullet(
    document,
    text,
    indent=0.25
):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.left_indent = (
        Inches(indent)
    )

    paragraph.paragraph_format.first_line_indent = (
        Inches(-0.14)
    )

    paragraph.paragraph_format.space_after = Pt(1.5)

    run = paragraph.add_run(
        "•  " + text
    )

    run.font.name = "Arial"
    run.font.size = Pt(10)


def add_left_right_line(
    document,
    left_text,
    right_text,
    left_bold=False,
    right_bold=False
):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(0)

    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.35),
        WD_TAB_ALIGNMENT.RIGHT
    )


    left = paragraph.add_run(
        left_text or ""
    )

    left.bold = left_bold

    left.font.name = "Arial"
    left.font.size = Pt(10)


    paragraph.add_run(
        "\t"
    )


    right = paragraph.add_run(
        right_text or ""
    )

    right.bold = right_bold

    right.font.name = "Arial"
    right.font.size = Pt(10)

    return paragraph


# =========================================================
# RESUME
# =========================================================

def resume_to_docx(
    resume: ResumeDocument
) -> BytesIO:

    document = Document()

    configure_document(
        document
    )


    # Shared template header
    add_candidate_header(
        document,
        resume.name,
        resume.contact
    )


    # PROFILE
    add_section_heading(
        document,
        "Profile"
    )


    for item in resume.profile:

        add_bullet(
            document,
            item
        )


    # EXPERIENCE
    add_section_heading(
        document,
        "Experience"
    )


    for position in resume.career_positions:

        add_left_right_line(
            document,
            position.job_title,
            position.dates,
            left_bold=True,
            right_bold=True
        )


        employer = add_left_right_line(
            document,
            position.employer,
            position.location
        )

        employer.paragraph_format.space_after = Pt(5)


    # FUNCTIONAL SKILL GROUPS
    for group in resume.skill_groups:

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(
            group.name
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)


        for accomplishment in group.accomplishments:

            add_bullet(
                document,
                accomplishment
            )


    # EDUCATION
    if resume.education:

        add_section_heading(
            document,
            "Education"
        )


        for item in resume.education:

            add_left_right_line(
                document,
                item.program_name,
                item.dates,
                left_bold=True,
                right_bold=True
            )


            school = add_left_right_line(
                document,
                item.school_name,
                item.location
            )

            school.paragraph_format.space_after = Pt(5)


    # ADDITIONAL TRAINING
    if resume.additional_training:

        add_section_heading(
            document,
            "Additional Training"
        )


        for item in resume.additional_training:

            add_left_right_line(
                document,
                item.course_name,
                item.date,
                left_bold=True,
                right_bold=True
            )


            provider = add_left_right_line(
                document,
                item.provider,
                item.location
            )

            provider.paragraph_format.space_after = Pt(5)


    output = BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output


# =========================================================
# COVER LETTER
# =========================================================

def cover_letter_to_docx(
    letter: CoverLetterDocument
) -> BytesIO:

    document = Document()

    configure_document(
        document
    )


    # Matching candidate header
    add_candidate_header(
        document,
        letter.candidate_name,
        letter.candidate_contact
    )


    # DATE
    if letter.date_line:

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(10)

        run = paragraph.add_run(
            letter.date_line
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)


    # RECIPIENT BLOCK
    recipient_lines = [

        letter.recipient.name,
        letter.recipient.position,
        letter.recipient.employer,
        letter.recipient.address,
        letter.recipient.phone,
        letter.recipient.email,
    ]


    recipient_lines = [
        line
        for line in recipient_lines
        if line
    ]


    for line in recipient_lines:

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(0)

        run = paragraph.add_run(
            line
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)


    if recipient_lines:

        paragraph.paragraph_format.space_after = Pt(10)


    # RE:
    if letter.subject_line:

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(10)

        run = paragraph.add_run(
            letter.subject_line
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)

        run.bold = True


    # SALUTATION
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(10)

    run = paragraph.add_run(
        letter.salutation
    )

    run.font.name = "Arial"
    run.font.size = Pt(10)


    # FOUR BODY PARAGRAPHS
    body = [

        letter.paragraph_1,
        letter.paragraph_2,
        letter.paragraph_3,
        letter.paragraph_4,
    ]


    for text in body:

        if not text:
            continue

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(10)

        run = paragraph.add_run(
            text
        )

        run.font.name = "Arial"
        run.font.size = Pt(10)


    # CLOSING
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(10)

    run = paragraph.add_run(
        letter.closing + ","
    )

    run.font.name = "Arial"
    run.font.size = Pt(10)


    paragraph = document.add_paragraph()

    run = paragraph.add_run(
        letter.candidate_name
    )

    run.font.name = "Arial"
    run.font.size = Pt(10)


    output = BytesIO()

    document.save(
        output
    )

    output.seek(0)

    return output
