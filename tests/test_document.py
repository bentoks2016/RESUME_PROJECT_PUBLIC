from docx import Document

from app.document import (
    cover_letter_to_docx,
    resume_to_docx,
)
from app.models import (
    CareerPosition,
    ContactInformation,
    CoverLetterDocument,
    EducationItem,
    RecipientInformation,
    ResumeDocument,
    SkillGroup,
    TrainingItem,
)


def document_text(stream):
    document = Document(stream)

    return "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text
    )


def sample_resume():
    return ResumeDocument(
        name="Alex Morgan",
        contact=ContactInformation(
            phone="+1 555 0100",
            email="alex@example.com",
            linkedin="linkedin.com/in/alex",
            location="Toronto, ON",
        ),
        profile=[
            (
                "Operations transformation professional "
                "with 11 years of manufacturing experience."
            ),
            "Led multi-site operational improvement.",
            "Bachelor of Business Administration.",
        ],
        career_positions=[
            CareerPosition(
                job_title="Transformation Manager",
                employer="Example Manufacturing",
                dates="2020-2025",
                location="Toronto, ON",
            )
        ],
        skill_groups=[
            SkillGroup(
                name="TRANSFORMATION & DELIVERY",
                accomplishments=[
                    "Led multi-site improvement.",
                    "Managed program risks.",
                ],
            )
        ],
        education=[
            EducationItem(
                program_name=(
                    "Bachelor of Business Administration"
                ),
                school_name="Example University",
                dates="2010-2014",
                location="Toronto, ON",
            )
        ],
        additional_training=[
            TrainingItem(
                course_name="Lean Six Sigma Green Belt",
                provider="Example Institute",
                date="2021",
                location="Toronto, ON",
            )
        ],
    )


def sample_cover_letter():
    return CoverLetterDocument(
        candidate_name="Alex Morgan",
        candidate_contact=ContactInformation(
            phone="+1 555 0100",
            email="alex@example.com",
            linkedin="linkedin.com/in/alex",
            location="Toronto, ON",
        ),
        date_line="September 16, 2026",
        recipient=RecipientInformation(
            name="Hiring Manager",
            position="Talent Acquisition",
            employer="Example Manufacturing",
            address="100 Example Street",
        ),
        subject_line=(
            "RE: Senior Operations Transformation Manager"
        ),
        salutation="Dear Hiring Manager,",
        paragraph_1="First substantive paragraph.",
        paragraph_2="Second substantive paragraph.",
        paragraph_3="Third substantive paragraph.",
        paragraph_4="Fourth substantive paragraph.",
        closing="Sincerely",
    )


def test_resume_docx_contains_candidate_header():
    text = document_text(
        resume_to_docx(sample_resume())
    )

    assert "Alex Morgan" in text
    assert "+1 555 0100" in text
    assert "alex@example.com" in text
    assert "linkedin.com/in/alex" in text
    assert "Toronto, ON" in text


def test_resume_docx_preserves_core_sections():
    text = document_text(
        resume_to_docx(sample_resume())
    )

    assert "PROFILE" in text
    assert "EXPERIENCE" in text
    assert "EDUCATION" in text
    assert "ADDITIONAL TRAINING" in text


def test_resume_docx_preserves_employment_fields():
    text = document_text(
        resume_to_docx(sample_resume())
    )

    assert "Transformation Manager" in text
    assert "Example Manufacturing" in text
    assert "2020-2025" in text
    assert "Toronto, ON" in text


def test_resume_docx_preserves_skill_group_content():
    text = document_text(
        resume_to_docx(sample_resume())
    )

    assert "TRANSFORMATION & DELIVERY" in text
    assert "Led multi-site improvement." in text
    assert "Managed program risks." in text


def test_resume_docx_preserves_education_and_training():
    text = document_text(
        resume_to_docx(sample_resume())
    )

    assert "Bachelor of Business Administration" in text
    assert "Example University" in text
    assert "Lean Six Sigma Green Belt" in text
    assert "Example Institute" in text


def test_cover_letter_docx_preserves_recipient_and_subject():
    text = document_text(
        cover_letter_to_docx(sample_cover_letter())
    )

    assert "Hiring Manager" in text
    assert "Talent Acquisition" in text
    assert "Example Manufacturing" in text
    assert "100 Example Street" in text
    assert (
        "RE: Senior Operations Transformation Manager"
        in text
    )
    assert "Dear Hiring Manager," in text


def test_cover_letter_docx_preserves_all_body_paragraphs():
    stream = cover_letter_to_docx(
        sample_cover_letter()
    )
    document = Document(stream)

    body = {
        paragraph.text
        for paragraph in document.paragraphs
    }

    for number in (
        "First",
        "Second",
        "Third",
        "Fourth",
    ):
        assert (
            f"{number} substantive paragraph."
            in body
        )


def test_cover_letter_docx_preserves_closing():
    text = document_text(
        cover_letter_to_docx(sample_cover_letter())
    )

    assert "Sincerely," in text
    assert "Alex Morgan" in text


def test_renderers_return_valid_docx_streams():
    for stream in (
        resume_to_docx(sample_resume()),
        cover_letter_to_docx(
            sample_cover_letter()
        ),
    ):
        assert stream.tell() == 0

        document = Document(stream)

        assert document.paragraphs


def test_optional_resume_sections_do_not_create_empty_headings():
    resume = sample_resume()
    resume.education = []
    resume.additional_training = []

    text = document_text(
        resume_to_docx(resume)
    )

    assert "EDUCATION" not in text
    assert "ADDITIONAL TRAINING" not in text


def test_optional_cover_letter_recipient_fields_may_be_blank():
    letter = sample_cover_letter()
    letter.recipient.position = ""
    letter.recipient.address = ""
    letter.recipient.phone = ""
    letter.recipient.email = ""

    text = document_text(
        cover_letter_to_docx(letter)
    )

    assert "Hiring Manager" in text
    assert "Example Manufacturing" in text
    assert "Dear Hiring Manager," in text


def test_safe_filename_removes_unsafe_characters():
    from app.main import safe_filename

    assert safe_filename(
        "../../Alex Morgan Resume?.docx"
    ) == "Alex_Morgan_Resume_docx"


def test_safe_filename_has_nonempty_fallback():
    from app.main import safe_filename

    assert safe_filename(" /// ") == "Document"


def test_download_resume_returns_docx():
    import json

    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)

    response = client.post(
        "/api/download/resume",
        data={
            "resume_json": json.dumps(
                sample_resume().model_dump()
            ),
            "filename": "Alex Morgan Resume",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    )
    assert (
        response.headers["content-disposition"]
        == 'attachment; filename="Alex_Morgan_Resume.docx"'
    )

    from io import BytesIO

    document = Document(
        BytesIO(response.content)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Alex Morgan" in text
    assert "PROFILE" in text
    assert "Transformation Manager" in text


def test_download_cover_letter_returns_docx():
    import json
    from io import BytesIO

    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)

    response = client.post(
        "/api/download/cover-letter",
        data={
            "cover_letter_json": json.dumps(
                sample_cover_letter().model_dump()
            ),
            "filename": "Alex Morgan Cover Letter",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    )
    assert response.headers["content-disposition"] == (
        'attachment; '
        'filename="Alex_Morgan_Cover_Letter.docx"'
    )

    document = Document(
        BytesIO(response.content)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Alex Morgan" in text
    assert "Dear Hiring Manager," in text
    assert "First substantive paragraph." in text
    assert "Fourth substantive paragraph." in text


def test_resume_uses_standard_document_configuration():
    stream = resume_to_docx(
        sample_resume()
    )
    document = Document(stream)

    section = document.sections[0]

    # python-docx stores dimensions in EMUs, so compare
    # their inch representations with a small tolerance.
    assert abs(
        section.top_margin.inches - 0.42
    ) < 0.01
    assert abs(
        section.bottom_margin.inches - 0.45
    ) < 0.01
    assert abs(
        section.left_margin.inches - 0.50
    ) < 0.01
    assert abs(
        section.right_margin.inches - 0.50
    ) < 0.01

    normal = document.styles["Normal"]

    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 10


def test_resume_and_cover_letter_share_document_configuration():
    resume_document = Document(
        resume_to_docx(sample_resume())
    )
    letter_document = Document(
        cover_letter_to_docx(
            sample_cover_letter()
        )
    )

    resume_section = resume_document.sections[0]
    letter_section = letter_document.sections[0]

    assert (
        resume_section.top_margin
        == letter_section.top_margin
    )
    assert (
        resume_section.bottom_margin
        == letter_section.bottom_margin
    )
    assert (
        resume_section.left_margin
        == letter_section.left_margin
    )
    assert (
        resume_section.right_margin
        == letter_section.right_margin
    )

    assert (
        resume_document.styles["Normal"].font.name
        == letter_document.styles["Normal"].font.name
        == "Arial"
    )


def test_safe_filename_removes_unsafe_characters():
    from app.main import safe_filename

    assert safe_filename(
        "../../Alex Morgan Resume?.docx"
    ) == "Alex_Morgan_Resume_docx"


def test_safe_filename_has_nonempty_fallback():
    from app.main import safe_filename

    assert safe_filename(" /// ") == "Document"


def test_download_resume_returns_docx():
    import json
    from io import BytesIO

    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)

    response = client.post(
        "/api/download/resume",
        data={
            "resume_json": json.dumps(
                sample_resume().model_dump()
            ),
            "filename": "Alex Morgan Resume",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="Alex_Morgan_Resume.docx"'
    )

    document = Document(
        BytesIO(response.content)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Alex Morgan" in text
    assert "PROFILE" in text
    assert "Transformation Manager" in text


def test_download_cover_letter_returns_docx():
    import json
    from io import BytesIO

    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)

    response = client.post(
        "/api/download/cover-letter",
        data={
            "cover_letter_json": json.dumps(
                sample_cover_letter().model_dump()
            ),
            "filename": "Alex Morgan Cover Letter",
        },
    )

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document"
    )
    assert response.headers["content-disposition"] == (
        'attachment; '
        'filename="Alex_Morgan_Cover_Letter.docx"'
    )

    document = Document(
        BytesIO(response.content)
    )

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    assert "Alex Morgan" in text
    assert "Dear Hiring Manager," in text
    assert "First substantive paragraph." in text
    assert "Fourth substantive paragraph." in text


def test_resume_uses_standard_document_configuration():
    document = Document(
        resume_to_docx(sample_resume())
    )

    section = document.sections[0]

    assert abs(
        section.top_margin.inches - 0.42
    ) < 0.01
    assert abs(
        section.bottom_margin.inches - 0.45
    ) < 0.01
    assert abs(
        section.left_margin.inches - 0.50
    ) < 0.01
    assert abs(
        section.right_margin.inches - 0.50
    ) < 0.01

    normal = document.styles["Normal"]

    assert normal.font.name == "Arial"
    assert normal.font.size.pt == 10


def test_resume_and_cover_letter_share_document_configuration():
    resume_document = Document(
        resume_to_docx(sample_resume())
    )
    letter_document = Document(
        cover_letter_to_docx(
            sample_cover_letter()
        )
    )

    resume_section = resume_document.sections[0]
    letter_section = letter_document.sections[0]

    assert (
        resume_section.top_margin
        == letter_section.top_margin
    )
    assert (
        resume_section.bottom_margin
        == letter_section.bottom_margin
    )
    assert (
        resume_section.left_margin
        == letter_section.left_margin
    )
    assert (
        resume_section.right_margin
        == letter_section.right_margin
    )

    assert (
        resume_document.styles["Normal"].font.name
        == letter_document.styles["Normal"].font.name
        == "Arial"
    )
